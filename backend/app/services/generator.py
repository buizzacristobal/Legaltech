"""Render lawsuit blocks + liquidation schedule into a .docx (Chilean judicial format:
Times New Roman 12, 1.5 line spacing, justified)."""
from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt, RGBColor

from app.engine.legal_templates import (
    Block, build_lawsuit, fmt_clp, fmt_date, fmt_pct, fmt_uf,
)
from app.schemas.instrument import Attorney, Instrument
from app.schemas.liquidation import Currency, LiquidationLedger

FONT = "Times New Roman"


def _style(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", FONT)
    normal.paragraph_format.line_spacing = 1.5
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(3)
        s.top_margin = s.bottom_margin = Cm(2.5)


def _para(doc: Document, text: str, *, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
          color: RGBColor | None = None):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    return p


def _render_blocks(doc: Document, blocks: list[Block]) -> None:
    for b in blocks:
        if b.kind == "banner":
            _para(doc, b.text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  color=RGBColor(0xC0, 0x00, 0x00))
        elif b.kind == "suma":
            _para(doc, b.text, bold=True)
        elif b.kind == "heading":
            _para(doc, b.text, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        elif b.kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            _para(doc, b.text)


def _schedule(doc: Document, ledger: LiquidationLedger) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _para(doc, "ANEXO: LIQUIDACIÓN DE LA DEUDA", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    p = ledger.params
    is_uf = p.currency is Currency.UF
    money = fmt_uf if is_uf else fmt_clp
    _para(doc, f"Capital: {money(p.principal)}. Emisión: {fmt_date(p.issue_date)}. "
               f"Vencimiento: {fmt_date(p.maturity_date)}. Corte: {fmt_date(p.cutoff_date)}.")
    headers = ["Período", "Días", "Tasa anual", "Interés del mes", "Interés acumulado",
               "Total acumulado ($)"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, h in zip(table.rows[0].cells, headers):
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for m in ledger.months:
        row = table.add_row().cells
        for cell, v in zip(row, [m.period, str(m.days), fmt_pct(m.annual_rate),
                                 money(m.interest), money(m.accrued_interest),
                                 fmt_clp(m.accumulated_total_clp)]):
            cell.text = v
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.line_spacing = 1.0
                for r in para.runs:
                    r.font.size = Pt(10)
                    r.font.name = FONT
    _para(doc, f"Capital: {fmt_clp(ledger.principal_clp)} — Intereses: "
               f"{fmt_clp(ledger.interest_clp)} — TOTAL: {fmt_clp(ledger.total_clp)}",
          bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    for w in ledger.warnings:
        _para(doc, f"Advertencia: {w}", align=WD_ALIGN_PARAGRAPH.LEFT,
              color=RGBColor(0xC0, 0x00, 0x00))


def generate_lawsuit_docx(instr: Instrument, ledger: LiquidationLedger,
                          attorney: Attorney) -> bytes:
    doc = Document()
    _style(doc)
    _render_blocks(doc, build_lawsuit(instr, ledger, attorney))
    _schedule(doc, ledger)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
