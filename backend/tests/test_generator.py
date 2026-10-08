from datetime import date
from decimal import Decimal as D
from io import BytesIO

import pytest
from docx import Document

from app.engine.calculator import calculate_day_by_day_ledger
from app.engine.legal_templates import build_lawsuit, fmt_clp, fmt_date, fmt_uf
from app.schemas.instrument import Attorney, CourtJurisdiction, Pagare
from app.schemas.liquidation import Currency, DayCount, LiquidationParams
from app.services.generator import generate_lawsuit_docx

PARTY_C = {"name": "Banco Ejemplo SpA", "rut": "76.086.428-5"}
PARTY_D = {"name": "Juan Pérez", "rut": "12.345.678-5"}
ATT = Attorney(name="María Soto", rut="11.111.111-1", address="Huérfanos 1117, Santiago")


def pagare(**kw):
    base = dict(amount=D(1_000_000), currency=Currency.CLP, issue_date=date(2024, 1, 1),
                maturity_date=date(2024, 1, 31), agreed_rate_annual=D(24),
                creditor=PARTY_C, debtor=PARTY_D,
                jurisdiction=CourtJurisdiction(court_name="12° Juzgado Civil", city="Santiago"))
    return Pagare(**{**base, **kw})


def ledger_for(p: Pagare, cutoff=date(2024, 2, 29)):
    from tests.test_calculator import TABLE
    return calculate_day_by_day_ledger(LiquidationParams(
        principal=p.amount, currency=p.currency, issue_date=p.issue_date,
        maturity_date=p.maturity_date, cutoff_date=cutoff,
        agreed_rate=p.agreed_rate_annual, day_count=DayCount.ACT_365), TABLE)


def text_of(data: bytes) -> str:
    d = Document(BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    parts += [c.text for t in d.tables for r in t.rows for c in r.cells]
    return "\n".join(parts)


def test_formatters():
    assert fmt_clp(D(1_023_836)) == "$1.023.836"
    assert fmt_uf(D("100.1370")) == "100,1370 UF"
    assert fmt_date(date(2024, 2, 9)) == "9 de febrero de 2024"


def test_docx_content_and_typography():
    p = pagare()
    led = ledger_for(p)  # 38,795 interest (see test_calculator)
    data = generate_lawsuit_docx(p, led, ATT)
    doc = Document(BytesIO(data))
    txt = text_of(data)
    for needle in ("BORRADOR", "S.J.L. EN LO CIVIL DE SANTIAGO", "HECHOS", "DERECHO",
                   "PETITORIO", "PRIMER OTROSÍ", "SEGUNDO OTROSÍ", "TERCER OTROSÍ: Patrocinio",
                   "$1.000.000", "$38.795", "$1.038.795", "ANEXO", "María Soto"):
        assert needle in txt, needle
    assert "Acredita personería" not in txt  # no representatives -> omitted
    normal = doc.styles["Normal"]
    assert normal.font.name == "Times New Roman" and normal.font.size.pt == 12
    assert normal.paragraph_format.line_spacing == 1.5
    assert len(doc.tables[0].rows) == 1 + len(led.months)


def test_personeria_otrosi_when_representatives_exist():
    p = pagare(creditor_representatives=[{**PARTY_D, "capacity": "gerente general"}])
    blocks = build_lawsuit(p, ledger_for(p), ATT)
    heads = [b.text for b in blocks if b.kind == "heading"]
    assert any(h.startswith("TERCER OTROSÍ: Acredita personería") for h in heads)
    assert any(h.startswith("CUARTO OTROSÍ: Patrocinio") for h in heads)


def test_jurisdiction_required():
    p = pagare(jurisdiction=None)
    with pytest.raises(ValueError):
        build_lawsuit(p, ledger_for(p), ATT)


def test_uf_instrument_states_both_currencies():
    p = pagare(amount=D(100), currency=Currency.UF, issue_date=date(2024, 2, 10),
               maturity_date=date(2024, 2, 10), agreed_rate_annual=None)
    led = ledger_for(p, cutoff=date(2024, 2, 20))
    txt = text_of(generate_lawsuit_docx(p, led, ATT))
    assert "100,0000 UF" in txt and "0,1370 UF" in txt
